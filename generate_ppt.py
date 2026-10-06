import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    DEEP_BLUE = RGBColor(30, 58, 138) # #1E3A8A
    BLUE = RGBColor(37, 99, 235)     # #2563EB
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)
    TEXT_DARK = RGBColor(30, 41, 59)
    TEXT_MUTED = RGBColor(100, 116, 139)
    GREEN = RGBColor(16, 185, 129)
    RED = RGBColor(239, 68, 68)

    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="國立東華大學 AI AGENT 實作總結"):
        # Header category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, NAVY)

    # Accent decorative element
    dec = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(1.2), Inches(0.08))
    dec.fill.solid()
    dec.fill.fore_color.rgb = BLUE
    dec.line.fill.background()

    # Subtitle Tag
    tag_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(10), Inches(0.5))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "NDHU CAMPUS AGENT LAB · PRACTICE REPO"
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = RGBColor(147, 197, 253)

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(2.0))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "讓 AI 幫你完成一件校園小事"
    p_t.font.size = Pt(40)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(255, 255, 255)

    p_t2 = tf_t.add_paragraph()
    p_t2.text = "Human-in-the-Loop 課堂實作成果與安全審查報告"
    p_t2.font.size = Pt(24)
    p_t2.font.color.rgb = RGBColor(203, 213, 225)
    p_t2.space_before = Pt(12)

    # Metadata card at bottom
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.4))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
    meta_box.line.color.rgb = RGBColor(51, 65, 85)

    tf_m = meta_box.text_frame
    tf_m.word_wrap = True
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "Repository: pinan12/agent-lab-w05  |  路線: individual 個人  |  工具: Antigravity AI Agent"
    p_m1.font.size = Pt(14)
    p_m1.font.bold = True
    p_m1.font.color.rgb = RGBColor(241, 245, 249)

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "完成了任務 A（檔案整理）、B（活動挑選器）、C（器材清理）、D（安全退回）與完整驗收紀錄"
    p_m2.font.size = Pt(13)
    p_m2.font.color.rgb = RGBColor(148, 163, 184)
    p_m2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 2: Core Philosophy & Four Judgments
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, LIGHT_BG)
    add_header(s2, "核心精神：培養人類審查者的四大判斷力")

    cards_data = [
        ("1. 看懂計畫", "Understand Plan", "動手前先檢查 AI 的讀取與產出範圍。確認其嚴守授權路徑，原檔唯讀，絕不未經同意刪檔。", BLUE),
        ("2. 親自驗收", "Inspect Results", "不盲信「已完成」宣告。透過 SHA-256 雜湊核對、邊界條件測試，親自檢查輸出檔案內容與規格是否吻合。", GREEN),
        ("3. 辨認確認點", "Recognize Checks", "辨識出哪些動作屬於不可逆高風險操作（例如：刪除重複檔、以 final2 認定定稿、自行猜測補齊缺失數據）。", RGBColor(217, 119, 6)),
        ("4. 堅決退回", "Reject & Propose", "當 AI 提出越權或具破壞性之計畫時，明確拒絕並提出可被接受的合規替代方案，主導整個工作流。", RED)
    ]

    card_width = Inches(2.75)
    card_gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (title, sub, desc, accent) in enumerate(cards_data):
        x = start_x + i * (card_width + card_gap)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), card_width, Inches(5.0))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR

        # Accent top bar
        bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.8), card_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.4)

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.color.rgb = accent
        p2.font.bold = True
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_DARK
        p3.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Task A (Club Files)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, LIGHT_BG)
    add_header(s3, "任務 A｜一張任務卡，整理社團檔案 (01-club-files)")

    # Left Box
    b_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = CARD_BG
    b_left.line.color.rgb = BORDER_COLOR
    tf_l = b_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.4)
    tf_l.margin_top = Inches(0.4)

    p = tf_l.paragraphs[0]
    p.text = "作業規範與嚴格限制"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_l = [
        "嚴格限制範圍：僅處理 input/ 資料夾中的 12 個文字檔案。",
        "原檔唯讀不動：原始檔案完全不修改、不刪除、不覆蓋。",
        "拒絕檔名猜測：檔名含有 final 或 final2 不等於定稿，皆為討論草案，均須完整保留。",
        "重複檔各留副本：公告與器材等內容 100% 相同之重複檔案，亦各自保留獨立副本供對照。"
    ]
    for it in items_l:
        p = tf_l.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(14)

    # Right Box
    b_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = CARD_BG
    b_right.line.color.rgb = BORDER_COLOR
    tf_r = b_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.4)
    tf_r.margin_top = Inches(0.4)

    p = tf_r.paragraphs[0]
    p.text = "分類結構與驗收成果"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_r = [
        "5 大分類目錄：\n  - 01-proposals（戶外30分/室內20分企劃、雨天備案）\n  - 02-planning（會議記錄、待辦清單、預算草案）\n  - 03-equipment（器材清單、器材備份）\n  - 04-publicity（公告、公告副本、海報標語）\n  - 05-feedback（活動問卷回饋）",
        "產出 manifest.json：12 筆物件對照來源、目的與理由。",
        "產出 report.md：詳細說明重複檔與不同版本分析報告。",
        "雜湊驗證：12 份副本與原檔之 SHA-256 雜湊 100% 完全相符。"
    ]
    for it in items_r:
        p = tf_r.add_paragraph()
        p.text = "✔ " + it
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: Task B (Campus Activity Picker)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, LIGHT_BG)
    add_header(s4, "任務 B｜單頁離線活動挑選器與雙語演進 (02-campus-picker)")

    b1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    b1.fill.solid()
    b1.fill.fore_color.rgb = CARD_BG
    b1.line.color.rgb = BORDER_COLOR
    tf1 = b1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.4)
    tf1.margin_top = Inches(0.4)

    p = tf1.paragraphs[0]
    p.text = "第一版核心功能 (B v1)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_b1 = [
        "單頁離線 HTML (index.html)：無 CDN、免聯網、雙擊即開。",
        "三維度條件篩選：地點（室內/室外/不限）、時間（15/30/60m，小於等於條件）、強度（低/中/不限）。",
        "無符合項目處理：嚴格顯示「沒有符合條件的活動」，不偷放寬條件。",
        "最近 5 次歷史紀錄：最新在最前，具「清除紀錄」按鈕。",
        "重設篩選：重設條件至不限/30/不限，但不清除歷史紀錄。"
    ]
    for it in items_b1:
        p = tf1.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    b2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    b2.fill.solid()
    b2.fill.fore_color.rgb = CARD_BG
    b2.line.color.rgb = BORDER_COLOR
    tf2 = b2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.4)
    tf2.margin_top = Inches(0.4)

    p = tf2.paragraphs[0]
    p.text = "實作修改：中英雙語對照模式 (B v2)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLUE

    items_b2 = [
        "痛點分析 (Before)：原先僅能單一切換純中文或純英文，無法同時對照雙語學習。",
        "需求提出 (Request)：增加「中英雙語 (Bilingual)」模式，同時呈現中英文標籤與活動名稱。",
        "實作成效 (After)：\n  - 右上角提供 [ 中文 | English | 中英雙語 ] 膠囊切換列。\n  - 抽中活動同時顯示中文主標題與英文副標題（如「整理書包 / Organize your bag」）。\n  - 歷史紀錄與篩選標籤均中英對照。",
        "修改性質：新需求（New requirement），已驗收通過。"
    ]
    for it in items_b2:
        p = tf2.add_paragraph()
        p.text = "★ " + it
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 5: Task C (Equipment Data Cleaning)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, LIGHT_BG)
    add_header(s5, "任務 C｜社團器材文字記錄清理 (03-equipment)")

    # 3 Stat Cards
    stat_data = [
        ("原始列數", "10", "包含全空資料列", NAVY),
        ("移除列數", "1", "第 6 列為全空物件 {}", RED),
        ("有效列數", "9", "成功輸出至 normalized.json", GREEN)
    ]
    for i, (st, num, sub, col) in enumerate(stat_data):
        sc = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(1.8), Inches(3.7), Inches(1.4))
        sc.fill.solid()
        sc.fill.fore_color.rgb = CARD_BG
        sc.line.color.rgb = BORDER_COLOR
        stf = sc.text_frame
        stf.word_wrap = True
        stf.margin_left = Inches(0.3)
        stf.margin_top = Inches(0.2)
        p = stf.paragraphs[0]
        p.text = st
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p2 = stf.add_paragraph()
        p2.text = num
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = col
        p3 = stf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_DARK

    # Bottom Details Card
    b_c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.5), Inches(11.733), Inches(3.3))
    b_c.fill.solid()
    b_c.fill.fore_color.rgb = CARD_BG
    b_c.line.color.rgb = BORDER_COLOR
    tf_c = b_c.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.4)
    tf_c.margin_top = Inches(0.3)

    p = tf_c.paragraphs[0]
    p.text = "資料品質異常分析報告 (issues.md 亮點)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY

    c_issues = [
        "數值規範守則：qty 欄位只能接受 0 或正整數。不補 0、不取絕對值、不猜測數量，缺失與異常原樣留存。",
        "缺失數量 (EQ04 紙張包)：qty 為空字串 \"\"，忠實保留原樣並列入報告待幹部核對。",
        "負數數量 (EQ05 膠帶)：qty 為 -1，保留負數數值並提出預警，不擅自修正為 1 或 0。",
        "資料衝突 (EQ02 延長線)：存在兩筆有效記錄，借出狀態皆為 borrowed，但數量分別為 2 與 3，發生衝突。全部保留供實物盤點。"
    ]
    for it in c_issues:
        p = tf_c.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 6: Task D (Safety Review & Rejection)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, LIGHT_BG)
    add_header(s6, "任務 D｜安全審查與退回不合理計畫 (04-review)")

    # Left: Bad plan
    b_bad = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    b_bad.fill.solid()
    b_bad.fill.fore_color.rgb = RGBColor(254, 242, 242) # light red
    b_bad.line.color.rgb = RGBColor(254, 202, 202)
    tf_bad = b_bad.text_frame
    tf_bad.word_wrap = True
    tf_bad.margin_left = Inches(0.4)
    tf_bad.margin_top = Inches(0.3)

    p = tf_bad.paragraphs[0]
    p.text = "❌ 模擬計畫中的五大致命問題 (bad-plan.txt)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RED

    bad_points = [
        "越權掃描全機：聲稱「把 Downloads 全部整理」，違反最小權限原則，恐存取私人敏感檔案。",
        "擅自刪除原檔：未經核准直接「刪除重複檔」，造成不可逆的檔案滅失風險。",
        "憑檔名妄斷定稿：把 final2 當最新版，漠視兩份企劃為完全不同活動之事實。",
        "捏造缺失數據：宣稱「找不到就補合理值」，破壞資料真實性（資料污染）。",
        "未經審查自動公開：未經人工驗收就「自動公開成果」，產生重大安全隱憂。"
    ]
    for pt in bad_points:
        p = tf_bad.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(153, 27, 27)
        p.space_before = Pt(8)

    # Right: Good plan
    b_good = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    b_good.fill.solid()
    b_good.fill.fore_color.rgb = RGBColor(240, 253, 244) # light green
    b_good.line.color.rgb = RGBColor(187, 247, 208)
    tf_good = b_good.text_frame
    tf_good.word_wrap = True
    tf_good.margin_left = Inches(0.4)
    tf_good.margin_top = Inches(0.3)

    p = tf_good.paragraphs[0]
    p.text = "✔ 退回聲明與五大替代原則 (my-rejection.md)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GREEN

    good_points = [
        "嚴格限制工作邊界：僅能讀取指定的 input 目錄，嚴禁跨出範圍觸碰其他資料夾。",
        "原始檔案絕對唯讀：原檔保持不動，成果輸出至 output 副本，重複檔案亦保留獨立副本。",
        "所有版本完整留存：所有歷史草案完整保留，由社團幹部開會比對決策。",
        "缺失異常原樣記錄：不瞎猜、不補值，異常欄位於問題報告中列出。",
        "成果本地存放驗收：僅輸出至本地資料夾，經由人類操作驗收確認後方可發布。"
    ]
    for pt in good_points:
        p = tf_good.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = RGBColor(22, 101, 52)
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 7: Verification Matrix (Table)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, LIGHT_BG)
    add_header(s7, "實作驗收測試矩陣 (Testing Verification Matrix)")

    rows, cols = 8, 4
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.733)
    height = Inches(4.8)

    table_shape = s7.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(1.6)
    table.columns[1].width = Inches(3.4)
    table.columns[2].width = Inches(4.7)
    table.columns[3].width = Inches(2.033)

    headers = ["任務", "測試情境 / 動作", "預期行為與觀察結果", "驗收狀態"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

    test_rows = [
        ("任務 A", "12 檔案雜湊驗證", "輸出之 12 份副本與原檔 SHA-256 雜湊 100% 吻合，原檔未被修改或刪除", "✅ 通過"),
        ("任務 B", "室外／15分鐘／中強度", "顯示「沒有符合條件的活動」，不偷放寬條件，歷史紀錄未增加", "✅ 通過"),
        ("任務 B", "室外／30分鐘／中強度", "唯一符合條件為 A09（在合適位置快走），每次抽取必定為 A09", "✅ 通過"),
        ("任務 B", "連續抽取 6 次", "歷史紀錄只保留最近 5 次成功抽樣，最新在最上方，舊紀錄依序退出", "✅ 通過"),
        ("任務 B", "重設篩選測試", "篩選條件重設回不限/30/不限，歷史紀錄完整保留不被清空", "✅ 通過"),
        ("任務 B", "中英雙語對照模式", "介面標籤、活動主標（中文）與副標（英文）皆雙語並列", "✅ 通過"),
        ("任務 C", "器材資料清理正規化", "10 列剔除 1 筆全空列，9 筆有效列保留，缺失與負數數量皆留存", "✅ 通過")
    ]

    for r_idx, r_data in enumerate(test_rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = TEXT_DARK
                if c_idx == 3:
                    p.font.bold = True
                    p.font.color.rgb = GREEN

    # -------------------------------------------------------------
    # SLIDE 8: Git Commit Timeline & Conclusion (Dark Theme)
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, NAVY)

    # Header
    cat_box = s8.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    p_cat = cat_box.text_frame.paragraphs[0]
    p_cat.text = "GIT WORKFLOW & CONCLUSION".upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = RGBColor(147, 197, 253)

    title_box = s8.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
    p_title = title_box.text_frame.paragraphs[0]
    p_title.text = "Git 提交歷程與人機協同總結"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(255, 255, 255)

    # Left: Commit list
    b_com = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    b_com.fill.solid()
    b_com.fill.fore_color.rgb = RGBColor(30, 41, 59)
    b_com.line.color.rgb = RGBColor(51, 65, 85)
    tf_com = b_com.text_frame
    tf_com.word_wrap = True
    tf_com.margin_left = Inches(0.3)
    tf_com.margin_top = Inches(0.3)

    p = tf_com.paragraphs[0]
    p.text = "完整 Git 提交歷程 (8 個 Commits)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(241, 245, 249)

    commits = [
        ("docs: add lab summary slides report", "加入簡報報告與總結文件"),
        ("record: learning record and screenshots", "完成實作紀錄與 evidence 總覽"),
        ("C: clean equipment records", "完成器材資料清理與異常問題報告"),
        ("D: rejection", "完成模擬不合理計畫之退回聲明"),
        ("B v2: add bilingual display mode", "挑選器新增中英雙語對照模式"),
        ("B v1: activity picker", "完成單頁活動挑選器基本功能"),
        ("A: organize club files", "完成社團檔案分類與 manifest"),
        ("Initial commit", "課堂初始模板倉庫建立")
    ]
    for c_msg, c_desc in commits:
        p = tf_com.add_paragraph()
        p.text = f"• {c_msg}\n   ↳ {c_desc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(203, 213, 225)
        p.space_before = Pt(4)

    # Right: Summary Box
    b_sum = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    b_sum.fill.solid()
    b_sum.fill.fore_color.rgb = RGBColor(30, 41, 59)
    b_sum.line.color.rgb = RGBColor(51, 65, 85)
    tf_sum = b_sum.text_frame
    tf_sum.word_wrap = True
    tf_sum.margin_left = Inches(0.4)
    tf_sum.margin_top = Inches(0.4)

    p = tf_sum.paragraphs[0]
    p.text = "人機協同實踐心得"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)

    sum_points = [
        "AI 是高效的副駕駛，人類掌握決策方向盤：AI 能在數秒內完成檔案分類、資料清洗與網頁撰寫，但唯有人類能給予嚴格的邊界規範與安全約束。",
        "驗收勝於宣稱：任何由 AI 執行的成果，都必須由人類親自核對雜湊值、親測極端條件，確保真實性。",
        "原子化提交的重要性：每做完一題就 Commit 與 Push，不僅讓進度有清晰的版本軌跡，更確保在公用電腦重開機時成果不遺失。"
    ]
    for sp in sum_points:
        p = tf_sum.add_paragraph()
        p.text = "💡 " + sp
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.space_before = Pt(14)

    output_path = r"c:\Users\NDHU CSIE\Downloads\10.06\NDHU_Agent_Lab_Summary.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
